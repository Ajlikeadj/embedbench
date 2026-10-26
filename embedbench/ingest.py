"""Document discovery and parsing for local folders."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from bs4 import BeautifulSoup

from embedbench.config import IngestionConfig
from embedbench.errors import IngestionError
from embedbench.logging_utils import get_logger
from embedbench.models import Document

logger = get_logger(__name__)

_MARKDOWN_TITLE_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", re.MULTILINE)
_HORIZONTAL_SPACE_RE = re.compile(r"[ \t\f\v]+")
_EXCESS_NEWLINES_RE = re.compile(r"\n{3,}")


@dataclass(slots=True)
class SkippedFile:
    """A file that was discovered but not ingested."""

    path: str
    reason: str

    def to_dict(self) -> dict[str, str]:
        return {"path": self.path, "reason": self.reason}


@dataclass(slots=True)
class IngestionResult:
    """Loaded documents plus diagnostics about skipped files."""

    root: Path
    documents: list[Document]
    skipped: list[SkippedFile] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(document.text) for document in self.documents)

    @property
    def doc_ids(self) -> set[str]:
        return {document.doc_id for document in self.documents}

    def to_dict(self) -> dict[str, Any]:
        return {
            "root": str(self.root),
            "documents": len(self.documents),
            "chars": self.total_chars,
            "skipped": [item.to_dict() for item in self.skipped],
        }


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if b"\x00" in data:
        raise ValueError("file appears to contain binary (NUL) data")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def clean_text(text: str) -> str:
    """Normalize whitespace while preserving paragraph and Markdown structure."""

    normalized = unicodedata.normalize("NFKC", text)
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    normalized = _HORIZONTAL_SPACE_RE.sub(" ", normalized)
    normalized = "\n".join(line.rstrip() for line in normalized.splitlines())
    normalized = _EXCESS_NEWLINES_RE.sub("\n\n", normalized)
    return normalized.strip()


def _markdown_title(text: str) -> str | None:
    match = _MARKDOWN_TITLE_RE.search(text)
    return match.group(1).strip() if match else None


def _load_text(path: Path) -> tuple[str, str | None, dict[str, Any]]:
    raw = path.read_bytes()
    text = clean_text(_decode_text(raw))
    title = _markdown_title(text) if path.suffix.lower() in {".md", ".markdown"} else None
    return text, title, {}


def _load_html(path: Path) -> tuple[str, str | None, dict[str, Any]]:
    raw = path.read_bytes()
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "template", "nav", "footer"]):
        tag.decompose()

    title: str | None = None
    if soup.title and soup.title.get_text(strip=True):
        title = soup.title.get_text(" ", strip=True)
    elif soup.find("h1"):
        title = soup.find("h1").get_text(" ", strip=True)

    text = clean_text(soup.get_text("\n"))
    return text, title, {"format": "html"}


def _load_pdf(path: Path) -> tuple[str, str | None, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise IngestionError("PDF support requires pypdf: pip install pypdf") from exc

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise ValueError("PDF is encrypted and requires a password")
            except (NotImplementedError, AttributeError) as exc:
                raise ValueError("PDF encryption is not supported") from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except IngestionError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not parse PDF: {exc}") from exc

    text = clean_text("\n\n".join(pages))
    title: str | None = None
    try:
        metadata = reader.metadata or {}
        raw_title = metadata.get("/Title")
        if raw_title:
            title = str(raw_title).strip() or None
    except (AttributeError, TypeError):  # pragma: no cover - malformed metadata
        title = None
    return text, title, {"pages": len(pages), "format": "pdf"}


def _load_docx(path: Path) -> tuple[str, str | None, dict[str, Any]]:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise IngestionError(
            "DOCX support requires python-docx: pip install python-docx"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ValueError(f"Could not parse DOCX: {exc}") from exc

    blocks: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                blocks.append(paragraph.text.strip())
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))

    text = clean_text("\n\n".join(blocks))
    title: str | None = None
    try:
        title = (document.core_properties.title or "").strip() or None
    except (AttributeError, ValueError):  # pragma: no cover - malformed properties
        title = None
    tables = len(document.tables)
    return text, title, {"format": "docx", "tables": tables}


_LOADERS: dict[str, Callable[[Path], tuple[str, str | None, dict[str, Any]]]] = {
    ".txt": _load_text,
    ".md": _load_text,
    ".markdown": _load_text,
    ".html": _load_html,
    ".htm": _load_html,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def _source_type(suffix: str) -> str:
    return {
        ".md": "markdown",
        ".markdown": "markdown",
        ".htm": "html",
        ".txt": "text",
    }.get(suffix, suffix.lstrip("."))


def _discover(
    root: Path,
    config: IngestionConfig,
) -> tuple[list[Path], list[SkippedFile]]:
    if not root.exists():
        raise IngestionError(f"Documents folder does not exist: {root}")
    if not root.is_dir():
        raise IngestionError(f"Documents path is not a folder: {root}")

    resolved_root = root.resolve()
    skip_dirs = {name.casefold() for name in config.skip_dirs}
    extensions = set(config.extensions)
    iterator = resolved_root.rglob("*") if config.recursive else resolved_root.glob("*")

    discovered: list[Path] = []
    skipped: list[SkippedFile] = []
    for path in iterator:
        try:
            if not path.is_file():
                continue
            relative = path.relative_to(resolved_root)
        except (OSError, ValueError):
            continue
        parts = relative.parts
        if any(part.startswith(".") for part in parts):
            continue
        if any(part.casefold() in skip_dirs for part in parts[:-1]):
            continue
        if path.suffix.lower() not in extensions:
            continue
        try:
            if path.is_symlink() and not path.resolve().is_relative_to(resolved_root):
                skipped.append(
                    SkippedFile(relative.as_posix(), "symlink points outside documents folder")
                )
                continue
            size = path.stat().st_size
        except OSError as exc:
            skipped.append(SkippedFile(relative.as_posix(), f"stat failed: {exc}"))
            continue
        if size > config.max_file_bytes:
            skipped.append(
                SkippedFile(
                    relative.as_posix(),
                    f"file is {size} bytes, above max_file_bytes={config.max_file_bytes}",
                )
            )
            continue
        discovered.append(path)

    discovered.sort(key=lambda item: item.relative_to(resolved_root).as_posix().casefold())
    return discovered, skipped


def load_document(path: Path, root: Path) -> Document:
    """Load and clean one document.  Raises ``ValueError`` for unreadable files."""

    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ValueError(f"Unsupported file type: {suffix or '<none>'}")
    text, title, metadata = loader(path)
    text = clean_text(text)
    if not text:
        raise ValueError("no extractable text (scanned PDFs and image-only files are unsupported)")
    relative = path.relative_to(root.resolve()).as_posix()
    metadata = dict(metadata)
    metadata["size_bytes"] = path.stat().st_size
    metadata["extension"] = suffix
    return Document(
        doc_id=relative,
        path=path,
        text=text,
        source_type=_source_type(suffix),
        title=title,
        metadata=metadata,
    )


def load_documents(root: Path, config: IngestionConfig) -> IngestionResult:
    """Discover and parse all supported documents below ``root``.

    A malformed file does not abort the run: it is reported in ``skipped`` and
    ingestion continues.  An empty corpus is an error.
    """

    root = Path(root)
    files, skipped = _discover(root, config)
    documents: list[Document] = []
    for path in files:
        try:
            document = load_document(path, root)
        except Exception as exc:  # a broken file should not lose the whole corpus
            relative = path.relative_to(root.resolve()).as_posix()
            logger.warning("Skipping %s: %s", relative, exc)
            skipped.append(SkippedFile(relative, str(exc)))
            continue
        documents.append(document)
        logger.debug(
            "Loaded %s (%d chars, title=%r)", document.doc_id, len(document.text), document.title
        )

    if not documents:
        details = ""
        if skipped:
            first = "; ".join(f"{item.path}: {item.reason}" for item in skipped[:3])
            details = f" Skipped files: {first}"
        raise IngestionError(
            f"No readable documents found in {root}. Supported extensions: "
            f"{', '.join(config.extensions)}.{details}"
        )

    logger.info(
        "Loaded %d documents (%d characters), skipped %d",
        len(documents),
        sum(len(document.text) for document in documents),
        len(skipped),
    )
    return IngestionResult(root=root, documents=documents, skipped=skipped)
