from __future__ import annotations

from pathlib import Path

import docx
import pytest

from embedbench.config import IngestionConfig
from embedbench.errors import IngestionError
from embedbench.ingest import clean_text, load_document, load_documents


def make_config(**overrides) -> IngestionConfig:
    values = {"max_file_bytes": 200_000}
    values.update(overrides)
    return IngestionConfig(**values)


def test_loads_markdown_html_and_text(tmp_path: Path) -> None:
    (tmp_path / "guide.md").write_text("# Install Guide\n\nRun the installer.\n", encoding="utf-8")
    (tmp_path / "notes.txt").write_text(
        "Plain text notes.\n\nSecond paragraph.\n", encoding="utf-8"
    )
    (tmp_path / "page.html").write_text(
        "<html><head><title>Security</title></head><body>"
        "<nav>Ignore me</nav><script>alert('x')</script>"
        "<h1>Security</h1><p>Encrypt everything.</p></body></html>",
        encoding="utf-8",
    )
    (tmp_path / "ignored.json").write_text("{}", encoding="utf-8")
    (tmp_path / ".hidden.md").write_text("# Hidden", encoding="utf-8")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "deep.txt").write_text("Nested content.", encoding="utf-8")

    result = load_documents(tmp_path, make_config())

    assert [document.doc_id for document in result.documents] == [
        "guide.md",
        "nested/deep.txt",
        "notes.txt",
        "page.html",
    ]
    by_id = {document.doc_id: document for document in result.documents}
    assert by_id["guide.md"].title == "Install Guide"
    assert by_id["guide.md"].source_type == "markdown"
    assert by_id["page.html"].title == "Security"
    assert "Ignore me" not in by_id["page.html"].text
    assert "alert" not in by_id["page.html"].text
    assert "Encrypt everything." in by_id["page.html"].text
    assert by_id["notes.txt"].source_type == "text"
    assert result.total_chars > 0


def test_non_recursive_ingestion_only_reads_top_level(tmp_path: Path) -> None:
    (tmp_path / "top.txt").write_text("top level", encoding="utf-8")
    nested = tmp_path / "sub"
    nested.mkdir()
    (nested / "deep.txt").write_text("deep", encoding="utf-8")

    result = load_documents(tmp_path, make_config(recursive=False))

    assert [document.doc_id for document in result.documents] == ["top.txt"]


def test_big_files_are_skipped_with_reason(tmp_path: Path) -> None:
    (tmp_path / "small.txt").write_text("small", encoding="utf-8")
    (tmp_path / "large.txt").write_text("x" * 2000, encoding="utf-8")

    result = load_documents(tmp_path, make_config(max_file_bytes=1024))

    assert [document.doc_id for document in result.documents] == ["small.txt"]
    assert result.skipped
    assert result.skipped[0].path == "large.txt"
    assert "max_file_bytes" in result.skipped[0].reason


def test_binary_text_file_is_skipped_not_fatal(tmp_path: Path) -> None:
    (tmp_path / "ok.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "binary.txt").write_bytes(b"\x00\x01\x02")
    result = load_documents(tmp_path, make_config())
    assert [document.doc_id for document in result.documents] == ["ok.txt"]
    assert any(item.path == "binary.txt" for item in result.skipped)


def test_empty_corpus_raises(tmp_path: Path) -> None:
    (tmp_path / "only.json").write_text("{}", encoding="utf-8")
    with pytest.raises(IngestionError, match="No readable documents"):
        load_documents(tmp_path, make_config())


def test_missing_folder_raises(tmp_path: Path) -> None:
    with pytest.raises(IngestionError, match="does not exist"):
        load_documents(tmp_path / "missing", make_config())


def test_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = docx.Document()
    document.core_properties.title = "Robot Manual"
    document.add_paragraph("The robot has four wheels.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Battery"
    table.rows[0].cells[1].text = "Eight hours"
    document.save(path)

    loaded = load_document(path, tmp_path)

    assert loaded.title == "Robot Manual"
    assert loaded.source_type == "docx"
    assert "four wheels" in loaded.text
    assert "Battery | Eight hours" in loaded.text
    assert loaded.metadata["tables"] == 1


def test_pdf_loader_uses_extracted_pages(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "manual.pdf"
    path.write_bytes(b"%PDF-1.4 fake")

    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        is_encrypted = False
        pages = [FakePage("Page one content."), FakePage("Page two content.")]
        metadata = {"/Title": "PDF Guide"}

    import pypdf

    monkeypatch.setattr(pypdf, "PdfReader", lambda _: FakeReader())

    loaded = load_document(path, tmp_path)

    assert loaded.title == "PDF Guide"
    assert "Page one content." in loaded.text
    assert "Page two content." in loaded.text
    assert loaded.metadata["pages"] == 2


def test_encrypted_pdf_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "secret.pdf"
    path.write_bytes(b"%PDF-1.4 fake")

    class FakeReader:
        is_encrypted = True

        def decrypt(self, password: str) -> int:
            return 0

    import pypdf

    monkeypatch.setattr(pypdf, "PdfReader", lambda _: FakeReader())

    with pytest.raises(ValueError, match="encrypted"):
        load_document(path, tmp_path)


def test_unsupported_extension_raises(tmp_path: Path) -> None:
    path = tmp_path / "data.bin"
    path.write_bytes(b"data")
    with pytest.raises(ValueError, match="Unsupported"):
        load_document(path, tmp_path)


def test_clean_text_normalizes_whitespace_and_preserves_paragraphs() -> None:
    raw = "Line one   with spaces\r\n\r\n\r\n\r\nLine two\tand tab\x00"
    cleaned = clean_text(raw)
    assert cleaned == "Line one with spaces\n\nLine two and tab"
